"""Generate DesignDocument.docx from ClaimsProcessor_DesignDocument.md."""
import re, sys, urllib.request
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Inches

ROOT = Path(__file__).resolve().parent
SRC = ROOT / 'ClaimsProcessor_DesignDocument.md'
OUT = ROOT / 'DesignDocument.docx'
IMG_DIR = ROOT / 'build' / 'diagrams'
IMG_DIR.mkdir(parents=True, exist_ok=True)


def add_page_number(paragraph):
    run = paragraph.add_run()
    for tag, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        if tag:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), tag); run._r.append(e)
        else:
            e = OxmlElement('w:instrText'); e.text = txt; run._r.append(e)


def shade_cell(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def add_horizontal_line(p):
    pPr = p._p.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    for k,v in (('w:val','single'),('w:sz','6'),('w:space','1'),('w:color','808080')): b.set(qn(k),v)
    pBdr.append(b); pPr.append(pBdr)


def paragraph_shade(p, color_hex):
    pPr = p._p.get_or_add_pPr(); shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),color_hex)
    pPr.append(shd)


def add_code_block(doc, text):
    p = doc.add_paragraph()
    paragraph_shade(p, 'F4F4F4')
    p.paragraph_format.left_indent = Cm(0.3); p.paragraph_format.right_indent = Cm(0.3)
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text); r.font.name = 'Consolas'; r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x1F,0x2A,0x44)
    pPr = p._p.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr')
    for edge in ('top','left','bottom','right'):
        b = OxmlElement(f'w:{edge}')
        for k,v in (('w:val','single'),('w:sz','4'),('w:space','4'),('w:color','CCCCCC')): b.set(qn(k),v)
        pBdr.append(b)
    pPr.append(pBdr)


def render_mermaid(src, idx):
    png = IMG_DIR / f'diagram_{idx:02d}.png'
    if png.exists() and png.stat().st_size > 0:
        return png
    print(f'  rendering mermaid #{idx}...')
    req = urllib.request.Request('https://kroki.io/mermaid/png', data=src.encode('utf-8'),
        headers={'Content-Type':'text/plain','User-Agent':'Mozilla/5.0'}, method='POST')
    png.write_bytes(urllib.request.urlopen(req, timeout=60).read())
    return png


def add_inline_runs(p, text):
    i = 0
    while i < len(text):
        if text.startswith('**', i):
            j = text.find('**', i+2)
            if j < 0: p.add_run(text[i:]); break
            r = p.add_run(text[i+2:j]); r.bold = True; i = j+2
        elif text.startswith('`', i):
            j = text.find('`', i+1)
            if j < 0: p.add_run(text[i:]); break
            r = p.add_run(text[i+1:j]); r.font.name='Consolas'; r.font.size=Pt(9)
            r.font.color.rgb = RGBColor(0xB0,0x10,0x40); i = j+1
        elif text.startswith('*', i) and (i == 0 or text[i-1] != '*'):
            j = text.find('*', i+1)
            if j < 0: p.add_run(text[i:]); break
            r = p.add_run(text[i+1:j]); r.italic = True; i = j+1
        else:
            nxt = len(text)
            for m in ('**','`','*'):
                k = text.find(m, i)
                if k >= 0 and k < nxt: nxt = k
            p.add_run(text[i:nxt]); i = nxt


def parse_blocks(md):
    lines = md.split('\n'); i = 0
    while i < len(lines):
        line = lines[i]
        m = re.match(r'^(#{1,4})\s+(.+?)\s*$', line)
        if m: yield ('h', (len(m.group(1)), m.group(2))); i+=1; continue
        if line.startswith('```'):
            lang = line[3:].strip(); start = i+1; j = start
            while j < len(lines) and not lines[j].startswith('```'): j += 1
            body = '\n'.join(lines[start:j])
            yield ('mermaid', body) if lang == 'mermaid' else ('code', body)
            i = j+1; continue
        if line.strip() == '---': yield ('hr', None); i+=1; continue
        if line.strip().startswith('|') and i+1 < len(lines) and re.match(r'^\s*\|[\s|:\-]+\|\s*$', lines[i+1]):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            yield ('table', (rows[0], rows[2:])); continue
        if line.startswith('> '):
            buf = []
            while i < len(lines) and lines[i].startswith('> '): buf.append(lines[i][2:]); i+=1
            yield ('quote', '\n'.join(buf)); continue
        if re.match(r'^\s*[-*]\s', line):
            buf = []
            while i < len(lines) and (re.match(r'^\s*[-*]\s', lines[i]) or (lines[i].startswith('  ') and buf)):
                buf.append(lines[i]); i += 1
            yield ('ul', buf); continue
        if re.match(r'^\s*\d+\.\s', line):
            buf = []
            while i < len(lines) and (re.match(r'^\s*\d+\.\s', lines[i]) or (lines[i].startswith('   ') and buf)):
                buf.append(lines[i]); i += 1
            yield ('ol', buf); continue
        if line.strip() == '': i += 1; continue
        buf = []
        while i < len(lines) and lines[i].strip() != '' and not lines[i].startswith('#') and not lines[i].startswith('```') and not lines[i].strip().startswith('|') and not lines[i].startswith('> ') and not re.match(r'^\s*[-*]\s', lines[i]) and not re.match(r'^\s*\d+\.\s', lines[i]):
            buf.append(lines[i]); i += 1
        yield ('p', ' '.join(buf).strip())


def main():
    md = SRC.read_text(encoding='utf-8')
    doc = Document()
    section = doc.sections[0]
    section.top_margin=Inches(0.7); section.bottom_margin=Inches(0.7)
    section.left_margin=Inches(0.9); section.right_margin=Inches(0.9)
    doc.styles['Normal'].font.name='Calibri'; doc.styles['Normal'].font.size=Pt(10.5)
    for lvl,sz,col in [(1,18,0x14365C),(2,14,0x1F4E79),(3,12,0x2E75B6),(4,11,0x4F81BD)]:
        s = doc.styles[f'Heading {lvl}']; s.font.name='Calibri'; s.font.size=Pt(sz)
        s.font.color.rgb = RGBColor((col>>16)&0xFF,(col>>8)&0xFF,col&0xFF); s.font.bold = (lvl <= 2)

    # Cover
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(180)
    r = p.add_run('Claims Transaction Processing System'); r.font.size=Pt(28); r.bold=True
    r.font.color.rgb = RGBColor(0x14,0x36,0x5C)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Design Document'); r.font.size=Pt(20); r.font.color.rgb = RGBColor(0x1F,0x4E,0x79)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Senior Software Engineer Assessment — ABC Health Insurance')
    r.font.size=Pt(13); r.italic=True; r.font.color.rgb = RGBColor(0x59,0x59,0x59)
    doc.add_paragraph().paragraph_format.space_after = Pt(120)
    for line in [
        'Java 25  ·  Spring Boot 3.4.5  ·  Apache POI  ·  OpenCSV  ·  Thymeleaf',
        'Spring Security (HTTP Basic)  ·  Micrometer + Prometheus  ·  springdoc-openapi',
        'JUnit 5  ·  Mockito  ·  JaCoCo (≥95%)  ·  Angular 18  ·  Bootstrap 5',
    ]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line); r.font.size=Pt(11); r.font.color.rgb=RGBColor(0x40,0x40,0x40)
    doc.add_paragraph().paragraph_format.space_after = Pt(200)
    for line, sz in [
        ('✓ Verified against SampleTransactionsProcessed: 0 / 19 mismatches', 12),
        ('✓ 80 / 80 unit + integration tests · JaCoCo line coverage ≥ 95%', 11),
        ('✓ Three runtime modes (Batch CLI · REST · Thymeleaf + Angular SPA)', 11),
    ]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line); r.font.size = Pt(sz); r.bold = (sz >= 12)
        r.font.color.rgb = RGBColor(0x1B,0x5E,0x20)

    fp = section.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(fp)
    r = fp.add_run('  ·  Claims Processor Design Document'); r.font.size=Pt(8); r.font.color.rgb=RGBColor(0x80,0x80,0x80)

    doc.add_page_break()

    midx = 0; first_h1 = True
    for kind, payload in parse_blocks(md):
        if kind == 'h':
            lvl, text = payload
            if lvl == 1 and first_h1: first_h1 = False; continue
            if lvl == 2 and text.lower().startswith('design document'): continue
            if lvl == 2 and re.match(r'^\d', text): doc.add_page_break()
            if lvl == 2 and text.lower().startswith('appendix'): doc.add_page_break()
            h = doc.add_heading(level=min(lvl,4)); add_inline_runs(h, text)
        elif kind == 'p':
            if not payload: continue
            p = doc.add_paragraph(); add_inline_runs(p, payload)
        elif kind == 'hr':
            add_horizontal_line(doc.add_paragraph())
        elif kind == 'code':
            add_code_block(doc, payload)
        elif kind == 'mermaid':
            midx += 1
            try:
                png = render_mermaid(payload, midx)
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(png), width=Cm(15.5))
            except Exception as e:
                print(f'  WARN: diagram {midx}: {e}', file=sys.stderr)
                add_code_block(doc, '[Mermaid source]\n' + payload)
        elif kind == 'table':
            header, rows = payload; n = len(header)
            tbl = doc.add_table(rows=1+len(rows), cols=n); tbl.style = 'Light Grid Accent 1'
            for c, h in enumerate(header):
                cell = tbl.rows[0].cells[c]; cell.text = ''
                p = cell.paragraphs[0]; add_inline_runs(p, h)
                for run in p.runs:
                    run.bold = True; run.font.color.rgb = RGBColor(0xFF,0xFF,0xFF); run.font.size = Pt(10)
                shade_cell(cell, '1F4E79')
            for ri, row in enumerate(rows, start=1):
                for c in range(n):
                    cell = tbl.rows[ri].cells[c]; cell.text=''
                    p = cell.paragraphs[0]; add_inline_runs(p, row[c] if c < len(row) else '')
                    for run in p.runs: run.font.size = Pt(9.5)
                    if ri % 2 == 0: shade_cell(cell, 'F2F7FC')
        elif kind == 'ul':
            for item in payload:
                m = re.match(r'^(\s*)[-*]\s+(.+?)\s*$', item)
                if m:
                    p = doc.add_paragraph(style='List Bullet'); add_inline_runs(p, m.group(2))
        elif kind == 'ol':
            for item in payload:
                m = re.match(r'^\s*\d+\.\s+(.+?)\s*$', item)
                if m:
                    p = doc.add_paragraph(style='List Number'); add_inline_runs(p, m.group(1))
        elif kind == 'quote':
            p = doc.add_paragraph(); paragraph_shade(p, 'EEF2F7')
            p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.right_indent = Cm(0.5)
            r = p.add_run(payload); r.italic = True; r.font.color.rgb = RGBColor(0x40,0x40,0x40)

    doc.save(OUT)
    print(f'wrote {OUT}  ({OUT.stat().st_size // 1024} KB)')


if __name__ == '__main__':
    main()
